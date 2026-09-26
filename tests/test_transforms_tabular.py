"""Tests for tabular-text + spreadsheet compression.

Covers detection (content_detector), the CSV→SmartCrusher bridge
(tabular_ingest), router wiring (content_router), and binary spreadsheet
ingestion (spreadsheet_ingest / compress_spreadsheet).
"""

from __future__ import annotations

import datetime
import importlib.util
import random

import pytest

from headroom.transforms.content_detector import (
    ContentType,
    DetectionResult,
    _is_md_separator,
    _looks_like_prose,
    _try_detect_delimited,
    _try_detect_markdown_table,
    detect_content_type,
)
from headroom.transforms.content_router import (
    CompressionStrategy,
    ContentRouter,
    ContentRouterConfig,
    _read_output_should_be_protected,
)
from headroom.transforms.tabular_ingest import (
    TabularCompressionResult,
    TabularCompressor,
    parse_csv,
    parse_fixed_width,
    parse_markdown_table,
    parse_tabular,
    to_records,
)

_HAS_OPENPYXL = importlib.util.find_spec("openpyxl") is not None


# Reusable fixtures ----------------------------------------------------------

CSV = "name,age,city\nAlice,30,NYC\nBob,25,LA\nCara,40,SF"
TSV = "id\tval\tnote\n1\ta\tx\n2\tb\ty\n3\tc\tz"
MARKDOWN = "| name | age |\n| --- | --- |\n| Alice | 30 |\n| Bob | 25 |\n| Cara | 40 |"


def _verbose_markdown(rows: int = 40) -> str:
    body = "\n".join(
        f"| user_{i} | {20 + i} | city_{i % 5} | active | engineering |" for i in range(rows)
    )
    return "| name | age | city | status | dept |\n| --- | --- | --- | --- | --- |\n" + body


# Detection ------------------------------------------------------------------


@pytest.mark.parametrize(
    "content,fmt",
    [(CSV, "csv"), (TSV, "csv"), (MARKDOWN, "markdown")],
)
def test_detects_tabular(content: str, fmt: str) -> None:
    result = detect_content_type(content)
    assert result.content_type is ContentType.TABULAR
    assert result.metadata.get("format") == fmt
    assert result.confidence >= 0.6


@pytest.mark.parametrize(
    "content,expected",
    [
        # Search output must not be stolen by tabular.
        (
            "src/main.py:42:def process():\nsrc/util.py:10:import os\nsrc/x.py:5:return 1",
            ContentType.SEARCH_RESULTS,
        ),
        # Build/log output stays a log.
        (
            "2026-01-01 INFO starting\n2026-01-01 WARN slow\n2026-01-01 ERROR boom",
            ContentType.BUILD_OUTPUT,
        ),
        # JSON arrays still go to the JSON path.
        ('[{"a": 1}, {"a": 2}, {"a": 3}]', ContentType.JSON_ARRAY),
        # Prose with incidental commas must NOT be tabular.
        (
            "Hello there, friend.\nThis is a sentence, yes.\nAnother line, ok.",
            ContentType.PLAIN_TEXT,
        ),
    ],
)
def test_does_not_misroute_to_tabular(content: str, expected: ContentType) -> None:
    assert detect_content_type(content).content_type is expected


# Detection: fixed-width command output (#3652) ------------------------------


def _ls_issue_payload() -> str:
    # The exact payload from issue #3652.
    rng = random.Random(1)
    rows = [
        f"-rw-r--r--  1 tejas staff {rng.randint(1000, 99999)} Sep {d} 09:{d:02d} file_{d}.py"
        for d in range(1, 60)
    ]
    return "total 480\n" + "\n".join(rows)


LS_MACOS = (
    "total 64\n"
    "drwxr-xr-x  12 tejas  staff    384 Sep 18 09:01 .\n"
    "drwxr-xr-x   5 tejas  staff    160 Sep 17 11:20 ..\n"
    "-rw-r--r--   1 tejas  staff   1834 Sep 18 09:01 README.md\n"
    "-rw-r--r--   1 tejas  staff  18611 Sep 18 09:01 setup.py\n"
    "drwxr-xr-x   8 tejas  staff    256 Sep 18 09:01 src"
)
KUBECTL = (
    "NAME                     READY   STATUS    RESTARTS   AGE\n"
    "api-7d9f8b6c4-2xkqp      1/1     Running   0          3d\n"
    "api-7d9f8b6c4-9wz7m      1/1     Running   0          3d\n"
    "worker-5c8d7f9b8-lq2vx   1/1     Running   2          5h\n"
    "redis-0                  1/1     Running   0          12d"
)
PS_AUX = (
    "USER       PID %CPU %MEM    VSZ   RSS TTY      STAT START   TIME COMMAND\n"
    "root         1  0.0  0.1 168100 11520 ?        Ss   Sep17   0:04 /sbin/init\n"
    "root         2  0.0  0.0      0     0 ?        S    Sep17   0:00 [kthreadd]\n"
    "tejas     4121  1.2  2.3 912344 190220 pts/0  Sl+  09:01   0:12 python app.py\n"
    "tejas     4188  0.0  0.0  10072  3300 pts/1    R+   09:05   0:00 ps aux"
)
DF_H = (
    "Filesystem      Size  Used Avail Use% Mounted on\n"
    "/dev/nvme0n1p2  468G  201G  244G  46% /\n"
    "tmpfs            16G  1.2M   16G   1% /dev/shm\n"
    "/dev/nvme0n1p1  511M  6.1M  505M   2% /boot/efi\n"
    "tmpfs           3.2G  2.4M  3.2G   1% /run/user/1000"
)


@pytest.mark.parametrize(
    "content",
    [_ls_issue_payload(), LS_MACOS, KUBECTL, PS_AUX, DF_H],
    ids=["ls_issue", "ls_macos", "kubectl", "ps_aux", "df_h"],
)
def test_detects_fixed_width_command_output(content: str) -> None:
    result = detect_content_type(content)
    assert result.content_type is ContentType.TABULAR
    assert result.metadata["format"] == "fixed_width"
    assert result.metadata["columns"] >= 3


@pytest.mark.parametrize(
    "content,expected",
    [
        pytest.param(
            "- first item in the list\n- second item in the list\n- third item in the list\n- fourth item in the list",
            ContentType.PLAIN_TEXT,
            id="bullets",
        ),
        pytest.param(
            "1. install the package\n2. run the proxy\n3. wrap the agent\n4. check the stats",
            ContentType.PLAIN_TEXT,
            id="numbered",
        ),
        pytest.param(
            "SELECT id, name\nFROM users\nWHERE active = 1\nORDER BY name;\n-- comment\nLIMIT 10;",
            ContentType.PLAIN_TEXT,
            id="sql",
        ),
        pytest.param(
            "#define FOO 1\n#define BAR 2\n#define BAZ 3\n#define QUX 4",
            ContentType.PLAIN_TEXT,
            id="c_defines",
        ),
        pytest.param(
            'On branch main\nChanges not staged for commit:\n  (use "git add <file>..." to update what will be committed)\n'
            + "\n".join(f"\tmodified:   src/m_{i}.py" for i in range(10)),
            ContentType.PLAIN_TEXT,
            id="git_status",
        ),
        pytest.param(
            "3aa5012 perf(memory/budget): precompute word sets once\nc81378c fix(grok): preserve xAI model context metadata\nb0c19a2 fix(security): reject unauthenticated public proxy binds\n871bbde fix(proxy): reject Anthropic batch operations on Copilot\na29162b fix(dashboard): separate rolling cache economics by owner",
            ContentType.PLAIN_TEXT,
            id="git_log",
        ),
        pytest.param(
            "Headroom compresses tool output before it reaches the model, which saves\ntokens on long agent sessions. The router picks a compressor per content\ntype, and plain prose goes to Kompress, an ML model that drops words it\npredicts the reader can do without. That is fine for prose and wrong for\nrecords, where every field matters to whatever command runs next, so the\ndetector has to tell the two apart before anything is dropped at all.",
            ContentType.PLAIN_TEXT,
            id="wrapped_prose",
        ),
    ],
)
def test_fixed_width_does_not_claim_non_tables(content: str, expected: ContentType) -> None:
    assert detect_content_type(content).content_type is expected


# Detection — edge branches --------------------------------------------------


def test_is_md_separator_needs_two_columns() -> None:
    assert _is_md_separator("| --- | --- |")
    assert not _is_md_separator("| --- |")  # single column is not a separator
    assert not _is_md_separator("| a | b |")  # cells must be dashes


def test_markdown_table_needs_multiple_columns() -> None:
    # Valid separator below, but the header is a single column -> not a table.
    assert _try_detect_markdown_table(["x|", "---|---", "y|"]) is None


def test_delimited_needs_three_rows() -> None:
    assert _try_detect_delimited(["a,b,c", "1,2,3"]) is None


def test_delimited_rejects_delimiter_only_in_header() -> None:
    # Header has commas but the data rows don't: no stable column count.
    assert _try_detect_delimited(["a,b,c", "plain", "text"]) is None


def test_delimited_rejects_inconsistent_columns() -> None:
    # Column count swings too much to be a real table.
    assert _try_detect_delimited(["a,b", "c,d", "e,f,g,h", "i,j,k,l,m"]) is None


def test_delimited_keeps_first_equal_confidence_delimiter() -> None:
    # Comma and semicolon are both consistent; the comma candidate is set first
    # and a later, no-better delimiter does not displace it.
    result = _try_detect_delimited(["a,b;c", "d,e;f", "g,h;i"])
    assert result is not None
    assert result.metadata["delimiter"] == ","


def test_looks_like_prose_distinguishes_sentences_from_rows() -> None:
    # Wordy cells (avg > 3 words/cell) read as prose even without end punctuation.
    assert _looks_like_prose(["the quick brown fox runs, over the lazy dog now"], ",")
    # Short field tuples are real CSV rows, not prose.
    assert not _looks_like_prose(["a,b,c", "1,2,3", "x,y,z"], ",")


# Parsers --------------------------------------------------------------------


def test_parse_csv_and_records() -> None:
    headers, rows = parse_csv(CSV)
    assert headers == ["name", "age", "city"]
    assert rows[0] == ["Alice", "30", "NYC"]
    records = to_records(headers, rows)
    assert records[1] == {"name": "Bob", "age": "25", "city": "LA"}


def test_parse_markdown_table_drops_separator() -> None:
    headers, rows = parse_markdown_table(MARKDOWN)
    assert headers == ["name", "age"]
    assert ["Alice", "30"] in rows
    assert all("---" not in cell for row in rows for cell in row)


def test_parse_tabular_rejects_ragged_fixed_width(monkeypatch) -> None:
    # Rows with differing cell counts can't be zipped under the headers
    # without misattributing columns (#1652) — must pass through.
    import headroom.transforms.tabular_ingest as ti

    monkeypatch.setattr(
        ti,
        "detect_content_type",
        lambda _c: DetectionResult(ContentType.TABULAR, 0.9, {"format": "fixed_width"}),
    )
    ragged = (
        "tool  installed  latest  status\n"
        "rtk  0.42.4  0.43.0  update available\n"
        "rtk  ✓  0.42.4  0.42.4  -  up-to-date"
    )
    assert ti.parse_tabular(ragged) is None


def test_parse_tabular_rejects_ragged_markdown(monkeypatch) -> None:
    import headroom.transforms.tabular_ingest as ti

    monkeypatch.setattr(
        ti,
        "detect_content_type",
        lambda _c: DetectionResult(ContentType.TABULAR, 0.9, {"format": "markdown"}),
    )
    ragged = "| a | b | c |\n| --- | --- | --- |\n| 1 | 2 | 3 |\n| 4 | 5 |"
    assert ti.parse_tabular(ragged) is None


def test_compress_passes_through_ragged_table(monkeypatch) -> None:
    import headroom.transforms.tabular_ingest as ti

    monkeypatch.setattr(
        ti,
        "detect_content_type",
        lambda _c: DetectionResult(ContentType.TABULAR, 0.9, {"format": "fixed_width"}),
    )
    ragged = (
        "tool  installed  latest  status\n"
        "rtk  0.42.4  0.43.0  update available\n"
        "rtk  ✓  0.42.4  0.42.4  -  up-to-date"
    )
    result = TabularCompressor().compress(ragged)
    assert not result.was_modified
    assert result.compressed == ragged


def _csv_with_an_oversized_cell() -> str:
    # csv.field_size_limit is 128 KB per cell; one pasted document, log excerpt
    # or base64 blob in a column goes past it.
    return "id,title,body\nl,short,ok\n2,long,{}\n".format("x" * 200_000)


def test_parse_csv_gives_up_on_a_cell_past_the_field_size_limit() -> None:
    headers, rows = parse_csv(_csv_with_an_oversized_cell())

    # csv.Error: field larger than field limit (131072) before this.
    assert (headers, rows) == ([], [])


def test_compress_passes_through_a_table_with_an_oversized_cell() -> None:
    content = _csv_with_an_oversized_cell()

    result = TabularCompressor().compress(content)

    assert not result.was_modified
    assert result.compressed == content


def test_parse_tabular_returns_none_for_non_tabular() -> None:
    assert parse_tabular("just a normal paragraph here") is None


def test_parse_fixed_width() -> None:
    headers, rows = parse_fixed_width("name    age   city\nAlice   30    NYC\nBob     25    LA")
    assert headers == ["name", "age", "city"]
    assert rows[0] == ["Alice", "30", "NYC"]


def test_to_records_empty_headers_returns_empty() -> None:
    assert to_records([], [["a", "b"]]) == []


def test_parse_csv_blank_returns_empty() -> None:
    assert parse_csv("   \n  \n") == ([], [])


def test_parse_markdown_table_too_short_returns_empty() -> None:
    assert parse_markdown_table("| only one row |") == ([], [])


def test_parse_fixed_width_too_short_returns_empty() -> None:
    assert parse_fixed_width("a single line") == ([], [])


def test_parse_tabular_dispatches_fixed_width(monkeypatch) -> None:
    # Drive the fixed_width dispatch branch directly with a stubbed detection
    # result, independent of the detector's thresholds.
    import headroom.transforms.tabular_ingest as ti

    monkeypatch.setattr(
        ti,
        "detect_content_type",
        lambda _c: DetectionResult(ContentType.TABULAR, 0.9, {"format": "fixed_width"}),
    )
    headers, rows, fmt = ti.parse_tabular("name    age\nAlice   30\nBob     25")
    assert fmt == "fixed_width"
    assert headers == ["name", "age"]
    assert rows[0] == ["Alice", "30"]


def test_parse_tabular_rejects_single_column_fixed_width(monkeypatch) -> None:
    import headroom.transforms.tabular_ingest as ti

    monkeypatch.setattr(
        ti,
        "detect_content_type",
        lambda _c: DetectionResult(ContentType.TABULAR, 0.9, {"format": "fixed_width"}),
    )
    # Single-space rows split into one cell each; that is not a table.
    assert (
        ti.parse_tabular("-rw-r--r-- 1 a b 1 f\n-rw-r--r-- 1 a b 2 g\n-rw-r--r-- 1 a b 3 h") is None
    )


def test_parse_tabular_none_when_no_data_rows_survive() -> None:
    # Detected as a markdown table, but it is header + separator rows only:
    # nothing survives as a data row, so parse_tabular bails to None.
    assert parse_tabular("| a | b |\n| --- | --- |\n| --- | --- |") is None


def test_compression_ratio_zero_for_empty_original() -> None:
    result = TabularCompressionResult(
        compressed="", original="", was_modified=False, fmt="csv", rows=0, columns=0
    )
    assert result.compression_ratio == 0.0


# Bridge compressor ----------------------------------------------------------


def test_verbose_markdown_compresses() -> None:
    result = TabularCompressor().compress(_verbose_markdown())
    assert result.was_modified
    assert len(result.compressed) < len(result.original)
    assert result.compression_ratio < 1.0
    assert result.fmt == "markdown"


def test_compact_unique_csv_passes_through() -> None:
    # All-unique compact rows have nothing losslessly removable.
    result = TabularCompressor().compress(CSV)
    assert not result.was_modified
    assert result.compressed == CSV


def test_non_tabular_passes_through_unmodified() -> None:
    # Unparseable prose returns the original content untouched.
    text = "just a normal paragraph here"
    result = TabularCompressor().compress(text)
    assert not result.was_modified
    assert result.compressed == text


# Router wiring --------------------------------------------------------------


def test_router_routes_tabular() -> None:
    result = ContentRouter().compress(_verbose_markdown())
    assert result.strategy_used is CompressionStrategy.TABULAR
    assert result.total_compressed_tokens <= result.total_original_tokens


def test_router_caches_tabular_compressor() -> None:
    router = ContentRouter()
    first = router._get_tabular_compressor()
    assert first is router._get_tabular_compressor()  # second call returns the cached instance


def test_router_tabular_passthrough_when_compressor_unavailable(monkeypatch) -> None:
    # Defensive guard: if the tabular compressor can't be constructed, routing to
    # TABULAR leaves content untouched instead of crashing.
    md = _verbose_markdown()
    router = ContentRouter()
    monkeypatch.setattr(router, "_get_tabular_compressor", lambda: None)
    result = router.compress(md)
    assert result.compressed == md
    assert result.tokens_saved == 0


def test_router_respects_disable_flag() -> None:
    # Disabling skips the tabular compressor: content passes through unchanged
    # (the selected strategy label may still read TABULAR, like other disabled
    # compressors).
    md = _verbose_markdown()
    cfg = ContentRouterConfig(enable_tabular_compressor=False)
    result = ContentRouter(cfg).compress(md)
    assert result.compressed == md
    assert result.tokens_saved == 0


# Router: tables never fall back to Kompress (#3652) -------------------------


def _record_kompress_calls(monkeypatch) -> list[str]:
    calls: list[str] = []

    def fake(self, content, context, question=None, target_ratio=None):
        calls.append(content)
        return "x", 1  # would "win" on savings if the router ever called it

    monkeypatch.setenv("HEADROOM_DETECT_BACKEND", "python")
    monkeypatch.setattr(ContentRouter, "_try_ml_compressor", fake)
    return calls


def test_router_keeps_ls_output_verbatim(monkeypatch) -> None:
    calls = _record_kompress_calls(monkeypatch)
    payload = _ls_issue_payload()
    result = ContentRouter(ContentRouterConfig()).compress(payload)
    assert result.compressed == payload
    assert calls == []
    assert result.strategy_used is CompressionStrategy.TABULAR


def test_router_does_not_kompress_a_ragged_csv(monkeypatch) -> None:
    calls = _record_kompress_calls(monkeypatch)
    csv = "id,name,city\n" + "\n".join(f"{i},user_{i},city_{i % 5}" for i in range(30))
    csv += "\n99,extra,field,here"
    assert detect_content_type(csv).content_type is ContentType.TABULAR
    result = ContentRouter(ContentRouterConfig()).compress(csv)
    assert result.compressed == csv
    assert calls == []


def test_fixed_width_read_stays_protected(monkeypatch) -> None:
    monkeypatch.setenv("HEADROOM_DETECT_BACKEND", "python")
    assert _read_output_should_be_protected(_ls_issue_payload()) is True
    csv = "id,name,city\n" + "\n".join(f"{i},user_{i},city_{i % 5}" for i in range(30))
    assert _read_output_should_be_protected(csv) is False


# Binary spreadsheet ingestion -----------------------------------------------


def test_rows_to_csv_drops_trailing_empty_rows_and_has_no_dangling_cr() -> None:
    """Trailing all-empty rows are dropped and the output has no stray ``\\r``.

    openpyxl's used-range routinely extends past the last data row, so a sheet
    commonly ends in ``(None, None, ...)`` tuples. Those were emitted as blank
    ``,`` rows, and ``csv.writer``'s default ``\\r\\n`` terminator combined with
    ``.strip("\\n")`` left a dangling ``\\r`` — noise fed straight to the LLM.
    """
    from headroom.transforms.spreadsheet_ingest import _rows_to_csv

    rendered = _rows_to_csv(
        [["Name", "Age"], ["Alice", "30"], [None, None], ["", "  "], [None, None]]
    )
    assert rendered == "Name,Age\nAlice,30"
    assert "\r" not in rendered

    # Interior empty rows are preserved (only the trailing run is dropped).
    assert _rows_to_csv([["a", "b"], [None, None], ["c", "d"], [None, None]]) == "a,b\n,\nc,d"

    # A fully empty sheet renders to the empty string.
    assert _rows_to_csv([[None, None], ["", ""]]) == ""


@pytest.mark.skipif(not _HAS_OPENPYXL, reason="openpyxl not installed")
def test_load_and_compress_xlsx(tmp_path) -> None:
    import openpyxl

    from headroom import compress_spreadsheet
    from headroom.transforms.spreadsheet_ingest import load_spreadsheet

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["id", "name", "dept", "status"])
    for i in range(40):
        ws.append([i, f"user_{i}", ["eng", "sales", "ops"][i % 3], "active"])
    wb.create_sheet("Empty")  # should be skipped
    path = tmp_path / "sample.xlsx"
    wb.save(path)

    sheets = load_spreadsheet(path)
    assert list(sheets) == ["Data"]
    assert sheets["Data"].splitlines()[0] == "id,name,dept,status"

    result = compress_spreadsheet(str(path))
    assert result.tokens_after <= result.tokens_before


@pytest.mark.skipif(not _HAS_OPENPYXL, reason="openpyxl not installed")
def test_compress_spreadsheet_empty_workbook_returns_empty(tmp_path) -> None:
    import openpyxl

    from headroom import compress_spreadsheet

    wb = openpyxl.Workbook()  # one empty sheet, no rows
    path = tmp_path / "empty.xlsx"
    wb.save(path)

    result = compress_spreadsheet(str(path))
    assert result.messages == []
    assert result.tokens_saved == 0


def test_load_xls_renders_cells_like_the_xlsx_loader(tmp_path) -> None:
    """xlrd hands back the raw storage, not the value.

    A date is the serial number Excel keeps it as, a boolean is 1 or 0, and
    every number is a double, so a whole number arrives as ``12.0``. The two
    loaders then disagree about the same workbook, and the date is no longer
    recoverable from the text.
    """
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")

    from headroom.transforms.spreadsheet_ingest import load_spreadsheet

    date_style = xlwt.XFStyle()
    date_style.num_format_str = "YYYY-MM-DD"

    book = xlwt.Workbook()
    sheet = book.add_sheet("Data")
    for column, heading in enumerate(["When", "Active", "Units", "Rate", "Text"]):
        sheet.write(0, column, heading)
    sheet.write(1, 0, datetime.date(2024, 1, 1), date_style)
    sheet.write(1, 1, True)
    sheet.write(1, 2, 12)
    sheet.write(1, 3, 1.5)
    sheet.write(1, 4, "ok")
    path = tmp_path / "legacy.xls"
    book.save(path)

    rows = load_spreadsheet(path)["Data"].splitlines()

    assert rows[0] == "When,Active,Units,Rate,Text"
    # 45292.0,1,12.0,1.5,ok before this.
    assert rows[1] == "2024-01-01 00:00:00,True,12,1.5,ok"


def test_load_xls_renders_a_time_only_cell_as_a_time(tmp_path) -> None:
    """A time carries no date, so xlrd reports year, month and day as zero."""
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    openpyxl = pytest.importorskip("openpyxl")

    from headroom.transforms.spreadsheet_ingest import load_spreadsheet

    time_style = xlwt.XFStyle()
    time_style.num_format_str = "HH:MM:SS"

    book = xlwt.Workbook()
    sheet = book.add_sheet("Data")
    sheet.write(0, 0, "Starts")
    sheet.write(1, 0, datetime.time(12, 0, 0), time_style)
    xls_path = tmp_path / "legacy.xls"
    book.save(xls_path)

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["Starts"])
    worksheet.append([datetime.time(12, 0, 0)])
    xlsx_path = tmp_path / "modern.xlsx"
    workbook.save(xlsx_path)

    # ValueError: year 0 is out of range before this.
    assert load_spreadsheet(xls_path)["Data"] == load_spreadsheet(xlsx_path)["Data"]
    assert load_spreadsheet(xls_path)["Data"].splitlines()[1] == "12:00:00"


def test_load_xls_and_xlsx_agree_above_the_exact_integer_range(tmp_path) -> None:
    """A double cannot hold consecutive integers past 2**53.

    ``_xls_cell`` converted any integral double with ``int()``, so a sheet
    holding 123456789012345678 rendered the double's exact value, 123456789012345680
    -- two fabricated digits presented to an agent as a precise identifier. The
    .xlsx loader has always rendered the float, which at least says
    "approximate", so bounding the conversion to the exactly-representable range
    keeps the ``12.0 -> 12`` fix from #3616 and restores agreement (#3695).
    """
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    openpyxl = pytest.importorskip("openpyxl")

    from headroom.transforms.spreadsheet_ingest import load_spreadsheet

    small, big = 12, 1.2345678901234568e17

    xls_book = xlwt.Workbook()
    xls_sheet = xls_book.add_sheet("Data")
    xls_sheet.write(0, 0, "Small")
    xls_sheet.write(0, 1, "Big")
    xls_sheet.write(1, 0, small)
    xls_sheet.write(1, 1, big)
    xls_path = tmp_path / "legacy.xls"
    xls_book.save(xls_path)

    # openpyxl is the reference the .xls path is written against, so the expected
    # rendering is the float repr it yields for the same value.
    openpyxl_wb = openpyxl.Workbook()
    openpyxl_sheet = openpyxl_wb.active
    openpyxl_sheet.title = "Data"
    openpyxl_sheet.append(["Small", "Big"])
    openpyxl_sheet.append([small, big])
    openpyxl_wb.save(tmp_path / "modern.xlsx")

    xls_row = load_spreadsheet(xls_path)["Data"].splitlines()[1]
    xlsx_row = load_spreadsheet(tmp_path / "modern.xlsx")["Data"].splitlines()[1]
    small_field, big_field = xls_row.split(",")

    # The #3616 win has to survive the bound: a small whole number is still an int.
    assert small_field == "12"
    # And the fabricated integer must be gone: the cell is rendered as the double
    # it is, which reads as an approximation instead of an exact identifier.
    assert big_field == repr(big)
    assert big_field != str(int(big))
    # Parity, asserted against the other loader rather than against my own
    # expectation. Small values agree verbatim; above the range openpyxl writes a
    # double with only 15 significant digits, so that side loses a digit on its
    # own and the rows cannot be string-equal. The promise this fix makes is
    # about magnitude: the two loaders agree to well within one unit in the last
    # place of the stored value (16 here), and neither hands the agent the
    # exact-looking decimal of the typed number.
    xlsx_small, xlsx_big = xlsx_row.split(",")
    assert small_field == xlsx_small == "12"
    assert abs(float(big_field) - float(xlsx_big)) <= 16


def test_load_xls_and_xlsx_agree_at_the_exact_integer_boundary(tmp_path) -> None:
    """2**53 and -2**53 are exactly representable and openpyxl loads them as
    integers, so the .xls path must convert them too - the bound is inclusive.
    One step outside, the double cannot hold the value; what matters is that no
    digits are invented, and the two loaders then differ only in the trailing
    ``.0`` that marks a value as approximate.
    """
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    openpyxl = pytest.importorskip("openpyxl")

    from headroom.transforms.spreadsheet_ingest import load_spreadsheet

    boundary = 2**53
    values = [boundary, -boundary, boundary - 2, boundary + 2]

    xls_book = xlwt.Workbook()
    xls_sheet = xls_book.add_sheet("Data")
    xls_sheet.write(0, 0, "Value")
    for row, value in enumerate(values, start=1):
        xls_sheet.write(row, 0, float(value))
    xls_path = tmp_path / "boundary.xls"
    xls_book.save(xls_path)

    xlsx_wb = openpyxl.Workbook()
    xlsx_sheet = xlsx_wb.active
    xlsx_sheet.title = "Data"
    xlsx_sheet.append(["Value"])
    for value in values:
        xlsx_sheet.append([int(value)])
    xlsx_path = tmp_path / "boundary.xlsx"
    xlsx_wb.save(xlsx_path)

    xls = [line.split(",")[0] for line in load_spreadsheet(xls_path)["Data"].splitlines()[1:]]
    xlsx = [line.split(",")[0] for line in load_spreadsheet(xlsx_path)["Data"].splitlines()[1:]]

    # Inside the range (and exactly on it) the two loaders agree verbatim.
    assert xls[0] == xlsx[0] == "9007199254740992"
    assert xls[1] == xlsx[1] == "-9007199254740992"
    assert xls[2] == xlsx[2] == "9007199254740990"
    # Above it the .xls side keeps the float marker, and the digits are the same.
    assert xls[3].removesuffix(".0") == xlsx[3] == "9007199254740994"
    assert all(not field.endswith(".0") or float(field) == int(float(field)) for field in xls)


def test_load_xls_and_xlsx_agree_on_the_same_values(tmp_path) -> None:
    """The reference: openpyxl is what the .xls path is matching."""
    xlwt = pytest.importorskip("xlwt")
    pytest.importorskip("xlrd")
    openpyxl = pytest.importorskip("openpyxl")

    from headroom.transforms.spreadsheet_ingest import load_spreadsheet

    date_style = xlwt.XFStyle()
    date_style.num_format_str = "YYYY-MM-DD"
    book = xlwt.Workbook()
    sheet = book.add_sheet("Data")
    sheet.write(0, 0, "When")
    sheet.write(0, 1, "Active")
    sheet.write(0, 2, "Units")
    sheet.write(1, 0, datetime.date(2024, 1, 1), date_style)
    sheet.write(1, 1, True)
    sheet.write(1, 2, 12)
    xls_path = tmp_path / "legacy.xls"
    book.save(xls_path)

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["When", "Active", "Units"])
    worksheet.append([datetime.date(2024, 1, 1), True, 12])
    xlsx_path = tmp_path / "modern.xlsx"
    workbook.save(xlsx_path)

    assert load_spreadsheet(xls_path) == load_spreadsheet(xlsx_path)


class _StubXlsCell:
    """The whole surface ``_xls_cell`` reads: xlrd's ``ctype`` and ``value``."""

    def __init__(self, ctype: int, value: object) -> None:
        self.ctype = ctype
        self.value = value


@pytest.mark.parametrize("value", [12.0, 1e15, float(2**53 - 1), float(2**53), float(-(2**53))])
def test_xls_cell_converts_exact_whole_numbers_to_int(value: float) -> None:
    """At or below 2**53 every integer is representable, so ``int()`` loses nothing.

    The bound is inclusive at both ends: +/-2**53 is exactly representable, and
    openpyxl reads the same value from an .xlsx as an ``int``, so excluding it
    would make the two loaders disagree at exactly the boundary.
    """
    xlrd = pytest.importorskip("xlrd")

    from headroom.transforms.spreadsheet_ingest import _xls_cell

    rendered = _xls_cell(_StubXlsCell(xlrd.XL_CELL_NUMBER, value), 0)

    assert isinstance(rendered, int)
    assert rendered == int(value)


@pytest.mark.parametrize(
    "value",
    [float(2**53 + 2), float(-(2**53) - 2), 1e16, 1e20, 123456789012345678.0],
)
def test_xls_cell_keeps_numbers_past_2_53_as_floats(value: float) -> None:
    """Past 2**53 ``int()`` would fabricate digits the workbook never held.

    ``2**53 + 2`` is the first whole number above the boundary (``2**53 + 1``
    is not representable at all), and ``-(2**53) - 2`` its negative mirror.

    xlrd hands back a double, and above 2**53 consecutive integers are no longer
    representable, so ``int()`` renders the double's exact value rather than the
    number that was typed: a cell holding 123456789012345678 prints as
    123456789012345680 -- an identifier that reads as exact and is wrong in its
    last two digits. The float repr says "approximate" out loud, and is also what
    the .xlsx loader shows for the same workbook.
    """
    xlrd = pytest.importorskip("xlrd")

    from headroom.transforms.spreadsheet_ingest import _xls_cell

    rendered = _xls_cell(_StubXlsCell(xlrd.XL_CELL_NUMBER, value), 0)

    assert isinstance(rendered, float)
    assert rendered == value


def test_load_spreadsheet_rejects_unknown_extension(tmp_path) -> None:
    from headroom.transforms.spreadsheet_ingest import load_spreadsheet

    bad = tmp_path / "data.txt"
    bad.write_text("a,b\n1,2\n")
    with pytest.raises(ValueError, match="Unsupported"):
        load_spreadsheet(bad)


def test_load_spreadsheet_missing_file(tmp_path) -> None:
    from headroom.transforms.spreadsheet_ingest import load_spreadsheet

    with pytest.raises(FileNotFoundError):
        load_spreadsheet(tmp_path / "nope.xlsx")
